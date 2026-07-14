import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session

from app import main
from app.config import Settings
from app.database import Base
from app.models import Signer, Template
from app.template_seed import prepare_beshtau_template_from_source


SOURCE = Path(__file__).resolve().parents[1] / "КП от Кристины.docx"


class BeshtauTemplateTest(unittest.TestCase):
    def test_replaces_hardcoded_signer_with_placeholders(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "beshtau_prepared.docx"
            prepare_beshtau_template_from_source(SOURCE, output)

            doc = Document(output)
            document_text = "\n".join(paragraph.text for paragraph in doc.paragraphs)

            self.assertIn("{{signer_title}}", document_text)
            self.assertIn("{{signer_name}}", document_text)
            self.assertNotIn("В.О. Галустян", document_text)

    def test_seeds_kuznetsov_and_makes_him_the_beshtau_default(self):
        engine = create_engine("sqlite://")
        Base.metadata.create_all(engine)
        with tempfile.TemporaryDirectory() as tmp, Session(engine) as db:
            previous_signer = Signer(title="Генеральный директор", name="В.О. Галустян", is_active=True)
            db.add(previous_signer)
            db.flush()
            template = Template(
                name="КП ООО «Бештау Электроникс»",
                organization="ООО «Бештау Электроникс»",
                default_signer_id=previous_signer.id,
            )
            db.add(template)
            db.commit()

            settings = Settings(
                seed_admin_email="",
                storage_dir=Path(tmp),
                beshtau_source_template_path=SOURCE,
                kartas_source_template_path=Path(tmp) / "missing-kartas.docx",
                nitrino_source_template_path=Path(tmp) / "missing-nitrino.docx",
            )
            with patch.object(main, "settings", settings):
                main.seed_initial_data(db)

            signer = db.scalar(select(Signer).where(Signer.name == "В.А. Кузнецов"))
            db.refresh(template)
            self.assertIsNotNone(signer)
            self.assertEqual(signer.title, "Генеральный директор")
            self.assertTrue(signer.is_active)
            self.assertEqual(template.default_signer_id, signer.id)


if __name__ == "__main__":
    unittest.main()
